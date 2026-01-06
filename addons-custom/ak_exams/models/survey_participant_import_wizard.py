# -*- coding: utf-8 -*-
import logging
import base64
import io
from odoo import models, fields, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

class SurveyParticipantImportWizard(models.TransientModel):
    _name = 'survey.participant.import.wizard'
    _description = 'Import Survey Participants from Excel'

    excel_file = fields.Binary(string=_('Excel File'), required=True,
                              help=_("Upload an Excel file to import participants (Rep and BM)."))
    excel_file_name = fields.Char(string=_('Excel File Name'))
    partner_id = fields.Many2one('res.partner', string=_('Company/Partner'), 
                                help=_("Select the company/partner to associate with the imported users."))
    tag_ids = fields.Many2many('res.partner.category', string=_('Tags'),
                              help=_("Tags to add to the imported users' partners."))
    
    analysis_result = fields.Html(string=_('Analysis Result'), readonly=True)
    state = fields.Selection([
        ('draft', 'Draft'),
        ('analyzed', 'Analyzed'),
        ('done', 'Done')
    ], default='draft', string='Status')

    def action_analyze(self):
        """Analyze the Excel file without importing"""
        self.ensure_one()
        if not self.excel_file:
            raise UserError(_("Please upload an Excel file first."))
        
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(_("The 'openpyxl' library is required."))

        try:
            decoded_file = base64.b64decode(self.excel_file)
            workbook = load_workbook(filename=io.BytesIO(decoded_file))
            sheet = workbook.active
            
            analysis_html = self._analyze_participants(sheet)
            self.write({
                'analysis_result': analysis_html,
                'state': 'analyzed'
            })
            
            return {
                'type': 'ir.actions.act_window',
                'res_model': 'survey.participant.import.wizard',
                'view_mode': 'form',
                'res_id': self.id,
                'target': 'new',
            }
        except Exception as e:
            _logger.exception("Error analyzing participants")
            raise UserError(_("Error processing Excel file: %s") % str(e))

    def action_import(self):
        self.ensure_one()
        if self.state != 'analyzed':
             raise UserError(_("Please analyze the file first."))
        
        if not self.excel_file:
            raise UserError(_("Please upload an Excel file first."))
        
        # Create automatic tag with timestamp
        timestamp = fields.Datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        auto_tag_name = f"Synchronized-{timestamp}"
        auto_tag = self.env['res.partner.category'].create({'name': auto_tag_name})
        
        # Add auto tag to selected tags
        self.tag_ids = [(4, auto_tag.id)]
        
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UserError(_("The 'openpyxl' library is required."))

        try:
            decoded_file = base64.b64decode(self.excel_file)
            workbook = load_workbook(filename=io.BytesIO(decoded_file))
            sheet = workbook.active
            
            stats = self._import_participants(sheet)
            
            message = _(
                "Import results:\n"
                "- Users Created: %(created)s\n"
                "- Users Updated: %(updated)s\n"
                "- Skipped: %(skipped)s\n"
                "- Teams Created: %(teams_created)s"
            ) % stats

            self.state = 'done'

            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {
                    'title': _('Import Successful'),
                    'message': message,
                    'sticky': False,
                    'type': 'success',
                }
            }
        except Exception as e:
            _logger.exception("Error importing participants")
            raise UserError(_("Error processing Excel file: %s") % str(e))

    def _find_user_by_email(self, email, external_id=None):
        """Find user by external_id first, then by email, trying both danone.com and nutricia.com domains"""
        User = self.env['res.users']
        
        # Try external_id first if provided
        if external_id:
            user = User.search([('external_id', '=', str(external_id))], limit=1)
            if user:
                return user
        
        # Try exact email
        user = User.search(['|', ('login', '=', email), ('email', '=', email)], limit=1)
        if user:
            return user
            
        # Try case-insensitive
        user = User.search(['|', ('login', '=ilike', email), ('email', '=ilike', email)], limit=1)
        if user:
            return user
            
        # Try alternative domain
        if '@danone.com' in email:
            alt_email = email.replace('@danone.com', '@nutricia.com')
        elif '@nutricia.com' in email:
            alt_email = email.replace('@nutricia.com', '@danone.com')
        else:
            return False
            
        user = User.search(['|', ('login', '=ilike', alt_email), ('email', '=ilike', alt_email)], limit=1)
        return user

    def _find_partner_by_email(self, email):
        """Find partner by email, trying both danone.com and nutricia.com domains"""
        Partner = self.env['res.partner']
        
        # Try exact email first
        partner = Partner.search([('email', '=ilike', email)], limit=1)
        if partner:
            return partner
            
        # Try alternative domain
        if '@danone.com' in email:
            alt_email = email.replace('@danone.com', '@nutricia.com')
        elif '@nutricia.com' in email:
            alt_email = email.replace('@nutricia.com', '@danone.com')
        else:
            return False
            
        partner = Partner.search([('email', '=ilike', alt_email)], limit=1)
        return partner

    def _analyze_participants(self, sheet):
        """Analyze the sheet and return HTML report"""
        report = []
        report.append("<table class='table table-bordered table-striped'>")
        report.append("<thead><tr><th>Row</th><th>Name</th><th>Email</th><th>Status</th><th>Details</th></tr></thead>")
        report.append("<tbody>")

        # Headers
        header = [str(cell.value).strip().upper() if cell.value else '' for cell in sheet[1]]
        
        header_mapping = {
            'FIRST&LAST NAME': 'name',
            'NAME': 'name',
            'AD SOYAD': 'name',
            'BUSINESS E-MAIL': 'email',
            'EMAIL': 'email',
            'E-POSTA': 'email',
            'DANONE ID': 'external_id',
            'ID': 'external_id',
            'TAKIM': 'team',
            'TEAM': 'team',
            'BÖLGE': 'region',
            'REGION': 'region',
        }

        col_map = {}
        for idx, h in enumerate(header):
            if h in header_mapping:
                col_map[header_mapping[h]] = idx
            else:
                for key, val in header_mapping.items():
                    if key in h:
                        col_map[val] = idx
                        break
        
        if 'email' not in col_map:
             return "<p class='text-danger'>Error: Excel file must contain an Email column.</p>"

        User = self.env['res.users']
        Partner = self.env['res.partner']
        Employee = self.env['hr.employee']
        Team = self.env['crm.team']

        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue

            email = row[col_map['email']] if 'email' in col_map and col_map['email'] < len(row) else None
            name = row[col_map['name']] if 'name' in col_map and col_map['name'] < len(row) else None
            external_id = row[col_map['external_id']] if 'external_id' in col_map and col_map['external_id'] < len(row) else None
            
            team_name = row[col_map['team']] if 'team' in col_map and col_map['team'] < len(row) else None
            region_name = row[col_map['region']] if 'region' in col_map and col_map['region'] < len(row) else None

            if not email:
                report.append(f"<tr><td>{row_idx}</td><td>{name or ''}</td><td>-</td><td><span class='badge badge-warning'>Skipped</span></td><td>Missing Email</td></tr>")
                continue
            
            email = str(email).strip().lower()
            status = ""
            details = []

            # Analyze user/partner status
            user = self._find_user_by_email(email, external_id)
            
            if user:
                if user.login != email:
                    status = "<span class='badge badge-warning'>Update Email</span>"
                    details.append(f"Found with different email: {user.login}")
                else:
                    status = "<span class='badge badge-info'>Update User</span>"
                details.append(f"User exists (ID: {user.id}, Login: {user.login})")
                
                # Check company parent update for existing user
                if self.partner_id and user.partner_id and user.partner_id.parent_id != self.partner_id:
                    details.append(f"Company parent will be updated to: {self.partner_id.name}")
                    
                # Check employee
                employee = Employee.search([('user_id', '=', user.id)], limit=1)
                if employee:
                    details.append(f"Employee exists (ID: {employee.id})")
                else:
                    details.append("Employee will be created")
            else:
                # Check for existing partner
                partner = self._find_partner_by_email(email)
                if not partner and name:
                    partner = Partner.search([('name', '=ilike', name)], limit=1)
                
                if partner:
                    status = "<span class='badge badge-warning'>Partner Exists</span>"
                    details.append(f"Partner exists (ID: {partner.id}) - Will create User linked to this Partner")
                    
                    # Check partner updates
                    if partner.email and partner.email.lower() != email:
                        details.append(f"Partner email will be updated: {partner.email} → {email}")
                    if self.partner_id and partner.parent_id != self.partner_id:
                        details.append(f"Company parent will be updated to: {self.partner_id.name}")
                else:
                    status = "<span class='badge badge-success'>Create User</span>"
                    details.append("New User will be created")
                    if name: details.append(f"Name: {name}")
                    if external_id: details.append(f"Ext ID: {external_id}")

            # Check teams
            if team_name:
                team_name = str(team_name).strip()
                if not Team.search([('name', '=', team_name), ('team_type', '=', 'G')], limit=1):
                    details.append(f"New Group Team: {team_name}")
            
            if region_name:
                region_name = str(region_name).strip()
                if not Team.search([('name', '=', region_name)], limit=1):
                    details.append(f"New Region Team: {region_name}")

            # Check external ID conflict
            if external_id:
                existing_user_ext = User.search([('external_id', '=', str(external_id)), ('login', '!=', email)], limit=1)
                if existing_user_ext:
                    status = "<span class='badge badge-danger'>Conflict</span>"
                    details.append(f"External ID {external_id} already used by {existing_user_ext.name} ({existing_user_ext.login})")

            report.append(f"<tr><td>{row_idx}</td><td>{name or ''}</td><td>{email}</td><td>{status}</td><td>{', '.join(details)}</td></tr>")

        report.append("</tbody></table>")
        return "".join(report)

    def _import_participants(self, sheet):
        stats = {
            'created': 0,
            'updated': 0,
            'skipped': 0,
            'teams_created': 0,
        }

        # Process all sheets in workbook
        workbook = sheet.parent
        for sheet in workbook.worksheets:
            _logger.info(f"Processing sheet: {sheet.title}")
            self._process_sheet(sheet, stats)
            
        return stats
    
    def _process_sheet(self, sheet, stats):
        """Process individual sheet based on sheet name"""
        
        # Skip empty sheets
        if sheet.max_row < 2:
            return
            
        # Headers
        header = [str(cell.value).strip().upper() if cell.value else '' for cell in sheet[1]]
        
        # Different header mappings for different sheets
        if sheet.title == 'SAHA':
            header_mapping = {
                'DANONE ID': 'external_id',
                'ID': 'external_id',
                'FIRST&LAST NAME': 'name',
                'NAME': 'name',
                'AD SOYAD': 'name',
                'BUSINESS E-MAIL': 'email',
                'EMAIL': 'email',
                'E-POSTA': 'email',
                'DEPARTMENT': 'department',
                'DEPARTMAN': 'department',
                'PHONE': 'phone',
                'TELEFON': 'phone',
                'TAKIM': 'team',
                'TEAM': 'team',
                'BÖLGE': 'region',
                'REGION': 'region',
            }
        else:  # BM sheet
            header_mapping = {
                'FIRST&LAST NAME': 'name',
                'NAME': 'name',
                'AD SOYAD': 'name',
                'BUSINESS E-MAIL': 'email',
                'EMAIL': 'email',
                'E-POSTA': 'email',
                'DEPARTMENT': 'department',
                'DEPARTMAN': 'department',
                'PHONE': 'phone',
                'TELEFON': 'phone',
                'BÖLGE': 'region',
                'REGION': 'region',
            }

        col_map = {}
        for idx, h in enumerate(header):
            if h in header_mapping:
                col_map[header_mapping[h]] = idx
            # Try partial match if exact match fails
            else:
                for key, val in header_mapping.items():
                    if key in h:
                        col_map[val] = idx
                        break
        
        if 'email' not in col_map:
             raise UserError(_("Excel file must contain an Email column."))

        User = self.env['res.users']
        Team = self.env['crm.team']
        
        # Cache teams to avoid repeated searches
        teams_cache = {} # name -> team_record

        _logger.info("Starting participant import...")
        
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_idx % 50 == 0:
                _logger.info(f"Processed {row_idx} rows...")
                self.env.cr.commit()  # Commit every 50 rows to avoid long transactions

            if not any(row):
                continue

            email = row[col_map['email']] if 'email' in col_map and col_map['email'] < len(row) else None
            if not email:
                stats['skipped'] += 1
                continue
            
            email = str(email).strip().lower()
            name = row[col_map['name']] if 'name' in col_map and col_map['name'] < len(row) else email.split('@')[0]
            if name:
                name = str(name).strip().upper()
            
            phone = row[col_map['phone']] if 'phone' in col_map and col_map['phone'] < len(row) else None
            external_id = row[col_map['external_id']] if 'external_id' in col_map and col_map['external_id'] < len(row) else None
            
            team_name = row[col_map['team']] if 'team' in col_map and col_map['team'] < len(row) else None
            region_name = row[col_map['region']] if 'region' in col_map and col_map['region'] < len(row) else None

            # Find or create user with domain switching
            user = self._find_user_by_email(email, external_id)

            # Update user email if different
            if user and user.login != email:
                user.write({'login': email, 'email': email})
                # Also update partner email if exists
                if user.partner_id and user.partner_id.email and user.partner_id.email.lower() != email:
                    user.partner_id.write({'email': email})
                _logger.info(f"Updated user and partner email to {email}")

            if not user:
                # Check if partner exists with domain switching
                partner = self._find_partner_by_email(email)
                if not partner and name:
                     partner = self.env['res.partner'].search([('name', '=ilike', name)], limit=1)
                
                user_vals = {
                    'name': name,
                    'login': email,
                    'email': email,
                    'lang': 'tr_TR',
                    'company_id': self.env.company.id,
                    'groups_id': [(6, 0, [self.env.ref('base.group_user').id, self.env.ref('ak_exams.group_nutricia_exams_participants').id])]
                }
                
                if partner:
                    # Update partner email if different
                    if partner.email and partner.email.lower() != email:
                        partner.write({'email': email})
                        _logger.info(f"Updated partner email from {partner.email} to {email}")
                    # Update partner parent if specified
                    if self.partner_id and partner.parent_id != self.partner_id:
                        partner.write({'parent_id': self.partner_id.id})
                    user_vals['partner_id'] = partner.id
                    user = User.create(user_vals)
                    stats['created'] += 1
                    _logger.info(f"Created user {email} linked to existing partner {partner.id}")
                else:
                    user = User.create(user_vals)
                    stats['created'] += 1
                    _logger.info(f"Created new user {email}")
            else:
                stats['updated'] += 1
                _logger.info(f"Updating existing user {email}")
            
            # Update user details
            vals = {'lang': 'tr_TR'}
            if name: vals['name'] = name
            if phone: vals['phone'] = phone
            if external_id: vals['external_id'] = str(external_id)
            
            # Ensure user has the participant group and is internal user
            group_participant = self.env.ref('ak_exams.group_nutricia_exams_participants')
            group_internal = self.env.ref('base.group_user')
            group_portal = self.env.ref('base.group_portal')
            group_public = self.env.ref('base.group_public')
            
            # Remove portal and public groups first
            groups_to_remove = []
            if group_portal.id in user.groups_id.ids:
                groups_to_remove.append(group_portal.id)
            if group_public.id in user.groups_id.ids:
                groups_to_remove.append(group_public.id)
                
            if groups_to_remove:
                vals['groups_id'] = [(3, gid) for gid in groups_to_remove]
                
            # Add required groups
            groups_to_add = []
            if group_participant.id not in user.groups_id.ids:
                groups_to_add.append(group_participant.id)
            if group_internal.id not in user.groups_id.ids:
                groups_to_add.append(group_internal.id)
                
            if groups_to_add:
                if 'groups_id' in vals:
                    vals['groups_id'].extend([(4, gid) for gid in groups_to_add])
                else:
                    vals['groups_id'] = [(4, gid) for gid in groups_to_add]
            
            # Update partner details
            if user.partner_id:
                partner_vals = {}
                if name and user.partner_id.name != name:
                    partner_vals['name'] = name
                if phone and user.partner_id.phone != phone:
                    partner_vals['phone'] = phone
                if partner_vals:
                    user.partner_id.write(partner_vals)

            if vals:
                user.write(vals)

            # Update partner parent if specified
            if self.partner_id and user.partner_id.parent_id != self.partner_id:
                user.partner_id.write({'parent_id': self.partner_id.id})
            
            # Add tags if specified
            if self.tag_ids:
                user.partner_id.write({'category_id': [(4, tag.id) for tag in self.tag_ids]})

            # Handle Teams based on sheet type
            group_team = False
            region_team = False
            
            if sheet.title == 'SAHA':
                # SAHA sheet: Create Group Team from Takım, Region Team from Bölge
                if team_name:
                    team_name = str(team_name).strip()
                    if team_name not in teams_cache:
                        group_team = Team.search([('name', '=', team_name), ('team_type', '=', 'G')], limit=1)
                        if not group_team:
                            group_team = Team.create({
                                'name': team_name,
                                'team_type': 'G',
                                'company_id': self.env.company.id
                            })
                            stats['teams_created'] += 1
                        teams_cache[team_name] = group_team
                    else:
                        group_team = teams_cache[team_name]
                        
                if region_name:
                    region_name = str(region_name).strip()
                    if region_name not in teams_cache:
                        region_team = Team.search([('name', '=', region_name)], limit=1)
                        if not region_team:
                            region_team = Team.create({
                                'name': region_name,
                                'team_type': 'R',
                                'parent_id': group_team.id if group_team else False,
                                'company_id': self.env.company.id
                            })
                            stats['teams_created'] += 1
                        elif group_team and region_team.parent_id != group_team:
                            region_team.write({'parent_id': group_team.id})
                        teams_cache[region_name] = region_team
                    else:
                        region_team = teams_cache[region_name]
                        if group_team and region_team.parent_id != group_team:
                            region_team.write({'parent_id': group_team.id})
            
            elif sheet.title == 'BM':
                # BM sheet: Only Region Team from Bölge, BM becomes team leader
                if region_name:
                    region_name = str(region_name).strip()
                    if region_name not in teams_cache:
                        region_team = Team.search([('name', '=', region_name)], limit=1)
                        if not region_team:
                            region_team = Team.create({
                                'name': region_name,
                                'team_type': 'R',
                                'company_id': self.env.company.id
                            })
                            stats['teams_created'] += 1
                        teams_cache[region_name] = region_team
                    else:
                        region_team = teams_cache[region_name]
                    
                    # Set BM as team leader
                    if user and region_team.user_id != user:
                        region_team.write({'user_id': user.id})

            # Assign user to team (Region Team)
            target_team = region_team
            if target_team:
                # Add user as member of the team
                if user.id not in target_team.member_ids.ids:
                    target_team.write({'member_ids': [(4, user.id)]})
                
                # Also set as sale_team_id on user if available (standard Odoo behavior often links these)
                if hasattr(user, 'sale_team_id'):
                    user.write({'sale_team_id': target_team.id})
            
            # Assign user to department based on sheet type
            if sheet.title == 'BM':
                # BM: Use Department column
                department_name = row[col_map['department']] if 'department' in col_map and col_map['department'] < len(row) else None
            else:
                # SAHA: Use Team column as department
                department_name = team_name if team_name else None
            
            # Determine job title and job position based on sheet name
            job_title = False
            job_position = None
            if sheet.title == 'SAHA':
                job_title = 'Rep'
                # Find or create Rep job position
                job_position = self.env['hr.job'].search([('name', '=', 'Rep')], limit=1)
                if not job_position:
                    job_position = self.env['hr.job'].create({'name': 'Rep'})
            elif sheet.title == 'BM':
                job_title = 'BM'
                # Find or create BM job position
                job_position = self.env['hr.job'].search([('name', '=', 'BM')], limit=1)
                if not job_position:
                    job_position = self.env['hr.job'].create({'name': 'BM'})
            
            if department_name:
                department_name = str(department_name).strip()
                # Search case-insensitive for existing department
                department = self.env['hr.department'].search([('name', '=ilike', department_name)], limit=1)
                if not department:
                    department = self.env['hr.department'].create({'name': department_name})
                elif department.name != department_name:
                    # Update department name if case is different
                    department.write({'name': department_name})
                
                # Check if user has an employee record
                employee = self.env['hr.employee'].search([('user_id', '=', user.id)], limit=1)
                if not employee:
                    employee_vals = {
                        'name': user.name,
                        'user_id': user.id,
                        'department_id': department.id,
                        'work_email': user.email,
                        'work_phone': user.phone,
                        'company_id': self.env.company.id,
                    }
                    if job_title:
                        employee_vals['job_title'] = job_title
                    if job_position:
                        employee_vals['job_id'] = job_position.id
                    employee = self.env['hr.employee'].create(employee_vals)
                else:
                    employee_vals = {
                        'department_id': department.id,
                        'company_id': self.env.company.id,
                    }
                    if job_title:
                        employee_vals['job_title'] = job_title
                    if job_position:
                        employee_vals['job_id'] = job_position.id
                    employee.write(employee_vals)
                    
                # Set as department manager for BM
                if sheet.title == 'BM':
                    department.write({'manager_id': employee.id})
            else:
                # Ensure employee exists even without department
                employee = self.env['hr.employee'].search([('user_id', '=', user.id)], limit=1)
                if not employee:
                    employee_vals = {
                        'name': user.name,
                        'user_id': user.id,
                        'work_email': user.email,
                        'work_phone': user.phone,
                        'company_id': self.env.company.id,
                    }
                    if job_title:
                        employee_vals['job_title'] = job_title
                    if job_position:
                        employee_vals['job_id'] = job_position.id
                    employee = self.env['hr.employee'].create(employee_vals)
                else:
                    employee_vals = {'company_id': self.env.company.id}
                    if job_title:
                        employee_vals['job_title'] = job_title
                    if job_position:
                        employee_vals['job_id'] = job_position.id
                    employee.write(employee_vals)

        return stats
